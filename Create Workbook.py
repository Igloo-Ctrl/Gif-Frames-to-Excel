from openpyxl import load_workbook, Workbook

amount_of_sheets = 57
excel_filename = "test.xlsx"


def create_work_book():
    wb = Workbook()
    wb.save(filename=excel_filename)
    for i in range(1, amount_of_sheets + 1):
        wb.create_sheet(title=f"Frame {i}")
    wb.save(filename=excel_filename)


if __name__ == '__main__':
    create_work_book()
